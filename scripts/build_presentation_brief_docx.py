from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "documents"
OUT = OUT_DIR / "unipath_max_presentation_brief.docx"


FONT = "Aptos"
BLACK = RGBColor(0, 0, 0)
BLUE = "D9EAF7"
DARK = "1F2937"
PALE = "F8FAFC"
GRID = "D9D9D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_borders(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text: str, bold: bool = False, color: RGBColor | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = color


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in [("Title", 22), ("Heading 1", 15), ("Heading 2", 12)]:
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)


def p(doc: Document, text: str = "", style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)
    return paragraph


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_status_table(doc: Document) -> None:
    p(doc, "Текущий статус реализации", "Heading 1")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for cell, text in zip(hdr.cells, ["Блок", "Статус", "Что важно показать"]):
        set_cell_shading(cell, DARK)
        set_cell_text(cell, text, bold=True, color=RGBColor(255, 255, 255))

    rows = [
        ("Student mini app", "Работает локально", "onboarding, Career GPS, opportunities, subscriptions, knowledge search"),
        ("Admin web panel", "Работает локально", "publish opportunity, knowledge management, analytics dashboard"),
        ("Backend API", "Проверен тестами", "health, JWT auth, tenant isolation, matching, notifications, OpenAPI, DATA-API smoke"),
        ("Docker", "Проверен агентом", "clean build with configurable ports and healthy services"),
        ("MAX mock mode", "Проверен", "notification is created once and marked simulated"),
        ("MAX real mode", "Нужна финальная проверка", "requires real bot token, HTTPS webhook, MAX_OPEN_APP_TARGET"),
    ]
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            set_cell_text(cell, text)
            if i % 2:
                set_cell_shading(cell, PALE)
    set_table_borders(table)
    table.autofit = True


def add_slide_table(doc: Document) -> None:
    p(doc, "Рекомендуемая структура презентации", "Heading 1")
    p(
        doc,
        "Презентация должна выглядеть как демонстрация работающего продукта, а не как рассказ о планах. "
        "Лучший ритм: проблема, пользовательский сценарий, техническое доказательство, статус готовности.",
    )
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for cell, text in zip(hdr.cells, ["Слайд", "Смысл", "Главный визуал", "Текстовый тезис"]):
        set_cell_shading(cell, DARK)
        set_cell_text(cell, text, bold=True, color=RGBColor(255, 255, 255))

    slides = [
        ("1", "Техническая информация для проверки", "compact table with links and credentials", "UniPath MAX is reproducible through Docker and has a verifiable API scenario."),
        ("2", "Проблема студента и университета", "split journey: student questions and admin content backlog", "Студенту нужна точная проверенная информация и понятный карьерный следующий шаг."),
        ("3", "Что такое UniPath MAX", "product map from MAX bot to mini app and admin panel", "Personal navigator from university questions to career opportunities."),
        ("4", "Student journey", "sequence of 5 product screens", "Onboarding data directly drives Career GPS and opportunity matching."),
        ("5", "Career GPS", "readiness score, gaps, next actions", "The value is deterministic skill-gap logic, not a black-box AI answer."),
        ("6", "Smart opportunities", "match card with percent, reasons, gaps, save button", "Every recommendation explains why it fits and what is missing."),
        ("7", "Verified knowledge", "search result with source, status, actual date, fallback", "The system does not invent answers when verified data is missing."),
        ("8", "Admin flow", "admin publishes opportunity and sees analytics", "University staff can manage verified content and trigger student value without DB access."),
        ("9", "MAX notification scenario", "event chain: publish -> match -> notification -> open app", "The bonus-worthy mechanic is proactive return through MAX after a relevant publication."),
        ("10", "Architecture and reliability", "simple modular monolith diagram and PASS checks", "Django REST API, PostgreSQL, React apps, service layer, Docker, tests, OpenAPI and smoke."),
        ("11", "Scale and rollout", "multi-university roadmap", "The architecture is tenant-ready and can scale from one university to several."),
        ("12", "Status and honest limitation", "green checklist plus one amber item", "Mock mode is verified; real MAX delivery needs token and public HTTPS rehearsal."),
    ]
    for i, row in enumerate(slides):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            set_cell_text(cell, text)
            if i % 2:
                set_cell_shading(cell, PALE)
    set_table_borders(table)


def add_screenshot_list(doc: Document) -> None:
    p(doc, "Скриншоты которые нужно подготовить", "Heading 1")
    bullets(
        doc,
        [
            "MAX bot entry point and mini app launch screen after real token is connected.",
            "Student Home or Onboarding with selected university, interests, skills and career goal.",
            "Career GPS with readiness score, strengths, gaps and next actions.",
            "Opportunities list with match percentage, reasons and gaps.",
            "Knowledge search for practice request with verified source and freshness fields.",
            "Knowledge fallback for an unknown query showing escalation instead of invented answer.",
            "Admin Opportunities screen before and after publishing a matching opportunity.",
            "Notification list or MAX message proving subscription delivery; in local mode the label must say simulated.",
            "Admin analytics dashboard with interaction counters.",
            "Terminal or API smoke result showing PASS for DATA-API scenario.",
        ],
    )


def add_demo_script(doc: Document) -> None:
    p(doc, "Демонстрационный сценарий для слайдов и защиты", "Heading 1")
    numbered(
        doc,
        [
            "Start the stack with seeded demo data.",
            "Log in as the demo student and show profile or onboarding data.",
            "Open Career GPS and point to skill gaps and next actions.",
            "Open Opportunities and show explainable match cards.",
            "Create or show a Backend subscription.",
            "Log in as admin and publish a verified Backend opportunity.",
            "Run or show the smoke check proving one notification was created without duplicates.",
            "If real MAX is connected, show the MAX message and return button opening the mini app.",
            "Search for practice in Knowledge and show verified source.",
            "Search for an impossible query and show safe fallback with escalation.",
        ],
    )


def add_messaging(doc: Document) -> None:
    p(doc, "Ключевые формулировки для слайдов", "Heading 1")
    bullets(
        doc,
        [
            "UniPath MAX helps a student understand what to do next: close skill gaps, find relevant opportunities and rely on verified university information.",
            "The system uses deterministic business logic for matching, Career GPS and knowledge fallback. It does not depend on an external LLM for core value.",
            "MAX is not only a container for the mini app. The strongest scenario is proactive notification when the university publishes a relevant opportunity.",
            "For local checks, MAX delivery is simulated explicitly and never presented as real delivery.",
            "For maximum scoring, the final rehearsal must prove real bot token, public HTTPS webhook and return from MAX message to the opportunity.",
        ],
    )


def add_visual_direction(doc: Document) -> None:
    p(doc, "Дизайн направление", "Heading 1")
    p(
        doc,
        "Рекомендуемый стиль: современный продуктовый deck для образовательной технологии. "
        "Визуально лучше держать интерфейс спокойным и доказательным: много реальных экранов, мало декоративных иллюстраций, "
        "четкие статусы, короткие подписи, крупные callout-цифры только там, где они подтверждены проверками.",
    )
    bullets(
        doc,
        [
            "Palette: white, near black, MAX/platform blue accents, restrained green for PASS and amber for not yet verified real MAX.",
            "Use product screenshots as primary visuals; avoid generic education stock images.",
            "Show one main user story across the whole deck so slides feel connected.",
            "Keep technical slides readable: architecture diagram, API endpoints and verification checklist should be compact.",
            "Do not claim real MAX platform bonus unless the live token and HTTPS delivery are actually tested.",
        ],
    )


def add_technical_appendix(doc: Document) -> None:
    doc.add_section(WD_SECTION.NEW_PAGE)
    p(doc, "Техническое приложение для дизайнера", "Heading 1")
    p(
        doc,
        "Эти детали можно использовать в последнем слайде или скрытом backup-слайде для проверки жюри.",
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for cell, text in zip(hdr.cells, ["Пункт", "Значение"]):
        set_cell_shading(cell, DARK)
        set_cell_text(cell, text, bold=True, color=RGBColor(255, 255, 255))
    rows = [
        ("Default student URL", "http://localhost:3000"),
        ("Default admin URL", "http://localhost:3001"),
        ("Default backend health", "http://localhost:8000/api/health/"),
        ("Port overrides", "BACKEND_PORT, STUDENT_PORT, ADMIN_PORT"),
        ("Student account", "student@demo.local / demo12345"),
        ("Admin account", "admin@demo.local / demo12345"),
        ("Editor account", "editor@demo.local / demo12345"),
        ("Smoke command", "backend/.venv/Scripts/python.exe backend/scripts/smoke_data_api.py --base-url http://localhost:8000"),
        ("Production MAX env", "MAX_BOT_TOKEN, MAX_WEBHOOK_SECRET, MAX_WEBHOOK_URL, MAX_OPEN_APP_TARGET, MAX_WEBAPP_BASE_URL"),
        ("Core endpoints", "/api/max/launch/, /api/student/career-gps, /api/student/opportunities, /api/admin/opportunities, /api/v1/notifications/"),
    ]
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            set_cell_text(cell, text)
            if i % 2:
                set_cell_shading(cell, PALE)
    set_table_borders(table)


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("UniPath MAX presentation brief")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle.add_run("Описание проекта для дизайнера презентации на основе текущей технической реализации").italic = True

    p(
        doc,
        "Этот документ нужен дизайнеру, чтобы собрать презентацию по уже работающему MVP UniPath MAX. "
        "Главная задача презентации - показать не набор экранов, а проверяемый продуктовый сценарий внутри MAX: "
        "студент проходит onboarding, получает карьерный маршрут, видит объяснимые возможности, подписывается на тему, "
        "а публикация новой возможности администратором приводит к уведомлению.",
    )
    p(
        doc,
        "В презентации важно честно разделить локально проверенный mock-режим MAX и финальный real-режим, "
        "который требует реального токена, публичного HTTPS и проверки доставки.",
    )

    p(doc, "Короткое описание продукта", "Heading 1")
    p(
        doc,
        "UniPath MAX - персональный цифровой навигатор студента от университетских вопросов до первой работы. "
        "Студент получает единый вход через MAX mini app: verified knowledge search, Career GPS, matching возможностей, "
        "подписки и уведомления. Университет получает admin web panel для публикации знаний и возможностей, управления карьерными ролями и просмотра аналитики.",
    )
    bullets(
        doc,
        [
            "Основная ценность работает без генеративной AI-зависимости.",
            "Matching и Career GPS объяснимы: score, reasons, gaps, next actions.",
            "Knowledge search отвечает только по verified data и показывает fallback, если надежного материала нет.",
            "MAX используется как точка входа и как канал возврата пользователя через уведомление.",
        ],
    )

    add_status_table(doc)
    add_slide_table(doc)
    add_screenshot_list(doc)
    add_demo_script(doc)
    add_messaging(doc)
    add_visual_direction(doc)
    add_technical_appendix(doc)

    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
        footer = section.footer.paragraphs[0]
        footer.text = "UniPath MAX presentation brief"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.name = FONT
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(100, 116, 139)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
